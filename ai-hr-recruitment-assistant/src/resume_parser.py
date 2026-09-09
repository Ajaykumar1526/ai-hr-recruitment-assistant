"""
Resume Parser
-------------
Extracts raw text from an uploaded PDF or DOCX resume, then asks the LLM
to convert that text into clean, structured JSON data.
"""

import json
import io
from typing import Optional

import fitz  # PyMuPDF
import docx
from anthropic import Anthropic

from src.config import ANTHROPIC_API_KEY, ANTHROPIC_MODEL, require_api_key


RESUME_SCHEMA_INSTRUCTIONS = """
You are an expert resume parser for an HR recruitment system.

Read the resume text below and extract the candidate's information into
STRICT JSON only (no markdown fences, no commentary, no extra text).

Use exactly this schema:
{
  "name": string,
  "email": string,
  "phone": string,
  "location": string,
  "summary": string,
  "education": [
    {"degree": string, "institution": string, "year": string, "score": string}
  ],
  "experience": [
    {"title": string, "company": string, "duration": string, "description": string}
  ],
  "skills": [string],
  "certifications": [string],
  "projects": [
    {"name": string, "description": string, "technologies": [string]}
  ],
  "total_experience_years": number
}

Rules:
- If a field is missing from the resume, use an empty string, empty list, or 0.
- "skills" should be a flat, deduplicated list of technical and soft skills.
- Keep descriptions concise (1-2 sentences).
- Return ONLY the JSON object, nothing else.
"""


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract plain text from a PDF file's bytes using PyMuPDF."""
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        for page in pdf:
            text_parts.append(page.get_text())
    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from a DOCX file's bytes using python-docx."""
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    # Also capture text inside tables (common in resumes with skill tables)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs).strip()


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Route to the correct extractor based on file extension."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX resume.")


def parse_resume_with_llm(resume_text: str, client: Optional[Anthropic] = None) -> dict:
    """Send extracted resume text to Claude and get back structured JSON."""
    require_api_key()
    if not resume_text.strip():
        raise ValueError("No readable text was found in this resume file.")

    client = client or Anthropic(api_key=ANTHROPIC_API_KEY)

    response = client.messages.create(
        model=ANTHROPIC_MODEL,
        max_tokens=2000,
        system=RESUME_SCHEMA_INSTRUCTIONS,
        messages=[{"role": "user", "content": f"Resume text:\n\n{resume_text[:12000]}"}],
    )

    raw_text = "".join(block.text for block in response.content if block.type == "text")
    return _safe_json_parse(raw_text)


def _safe_json_parse(raw_text: str) -> dict:
    """Strip accidental markdown fences and parse JSON, with a helpful error on failure."""
    cleaned = raw_text.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        cleaned = cleaned.replace("json\n", "", 1) if cleaned.startswith("json\n") else cleaned
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Could not parse structured data from the model response: {exc}")


def parse_resume(file_bytes: bytes, filename: str, client: Optional[Anthropic] = None) -> dict:
    """Full pipeline: file bytes -> extracted text -> structured candidate data."""
    text = extract_resume_text(file_bytes, filename)
    structured = parse_resume_with_llm(text, client=client)
    structured["_raw_text"] = text
    return structured
